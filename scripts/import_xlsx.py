#!/usr/bin/env python3
"""Convert the costed recipe workbook into the app's seed data.

Usage:
    python3 scripts/import_xlsx.py [path/to/workbook.xlsx] [path/to/seed.json]

Defaults to data/Jade_Seeliger_Recipes_Woolworths_Costed.xlsx and
src/data/seed.json. Needs openpyxl (`pip install openpyxl`).

The workbook is read with cached values, so open and save it in Excel once
after editing formulas so the values are up to date.
"""
import csv
import json
import re
import sys
from pathlib import Path

import openpyxl

ROOT = Path(__file__).resolve().parent.parent
DEFAULT_XLSX = ROOT / "data" / "Jade_Seeliger_Recipes_Woolworths_Costed.xlsx"
DEFAULT_OUT = ROOT / "src" / "data" / "seed.json"
NUTRITION_CSV = ROOT / "data" / "nutrition.csv"
SERVINGS_CSV = ROOT / "data" / "servings.csv"

# Shopping-list sections, assigned per product key. Anything not listed
# falls back to "Pantry".
SECTIONS = {
    "Fruit & veg": """artichokes baby_marrows baby_spinach beetroot brinjals onions brussels
        pumpkin slaw butternut celery chilli apples green_apples potatoes lemons carrots
        salad_leaves bananas berries iceberg pomegranate portabello red_onions avocado rosa_200
        rosa_400 rosa_600 red_pepper yellow_pepper sliced_mushrooms tomatoes spring_onions
        corn_cobs broccoli cauliflower spinach gem_squash chopped_spinach ginger garlic edamame
        tomato_onion_mix""",
    "Fresh herbs": "basil bay_leaves chives coriander mint parsley rosemary sage thyme rocket",
    "Meat, chicken & fish": """beef_shin bacon meatballs beef_strips chicken_portions chicken_thighs
        lamb_frikkadels chicken_thigh_fillets chicken_thighs_drums chicken_breast_fillets lean_mince
        rotisserie_chicken pork_fillet pork_sausages cooked_chicken_breast salami smoked_salmon""",
    "Dairy & eggs": """buttermilk plain_yoghurt goats_cheese strawberry_yoghurt eggs cheddar_grated
        mozzarella cottage_cheese cream_cheese feta butter parmesan fresh_cream white_cheddar milk
        low_fat_milk""",
    "Ready meals, soups & stocks": """fresh_chicken_soup beef_stock_liquid veg_stock mushroom_soup
        moroccan_soup tomato_soup grains_90s coconut_rice pesto""",
    "Bakery": "carrot_muffin roti hot_cross_buns bread",
    "Frozen": "frozen_veg haddock peas filo puff_pastry",
    "Baking": """baking_powder bicarb cake_flour self_raising_flour brown_sugar white_sugar corn_starch
        cocoa vanilla_essence desiccated_coconut choc_chips sprinkles evaporated_milk""",
    "Herbs, spices & stock": """beef_stock_powder chicken_stock_tubs chilli_powder cinnamon garam_masala
        allspice cumin italian_herbs mixed_spice oregano turmeric curry_paste
        brown_onion_thickener chicken_sauce_powder""",
    "Sauces, oils & condiments": """tomato_sauce hoisin rice_vinegar pickled_ginger avocado_oil bovril
        canola_oil chilli_oil capers english_mustard mayonnaise dijon olive_oil_spray german_mustard
        soy_sauce steak_sauce salsa chutney sesame_oil olive_oil balsamic_glaze truffle_oil wasabi
        wholegrain_mustard balsamic_vinegar lemon_olive_oil white_vinegar honey peanut_butter
        pickled_onions green_olives butter_chicken_sauce med_cook_in""",
    "Tins & jars": """chickpeas_tin corn_tin lentils_tin tuna tomato_puree passata pears_tin tinned_tomatoes
        kidney_beans coconut_milk_lite sardines tomato_paste cherry_tomatoes_tin""",
    "Pasta, grains & cereal": """egg_noodles bulgur dried_chickpeas oats all_bran corn_flakes orzo
        breadcrumbs ww_fusilli""",
    "Nuts, seeds & dried fruit": """mixed_seeds chia dates pecans pumpkin_seeds flaked_almonds mixed_nuts
        pistachios sunflower_seeds sesame_seeds dried_fruit""",
    "Snacks & treats": """milk_choc cadbury easter_eggs smarties biscoff cream_crackers pretzels popcorn""",
    "Wine": "chardonnay white_wine",
}
SECTION_OF = {key: section for section, keys in SECTIONS.items() for key in keys.split()}


def slug(text):
    return re.sub(r"[^a-z0-9]+", "-", text.lower()).strip("-")


def num(value):
    if value is None or value == "":
        return None
    return round(float(value), 4)


def main():
    xlsx = Path(sys.argv[1]) if len(sys.argv) > 1 else DEFAULT_XLSX
    out = Path(sys.argv[2]) if len(sys.argv) > 2 else DEFAULT_OUT
    wb = openpyxl.load_workbook(xlsx, data_only=True)

    # Products
    products = []
    ws = wb["Woolworths Products"]
    for row in ws.iter_rows(min_row=2):
        key = row[0].value
        if not key:
            continue
        link = row[10].hyperlink.target if row[10].hyperlink else None
        products.append({
            "key": key,
            "name": row[1].value,
            "today": num(row[2].value),
            "regular": num(row[3].value),
            "packSize": num(row[5].value),
            "unit": row[6].value,
            "note": row[9].value or "",
            "link": link or "",
            "section": SECTION_OF.get(key, "Pantry"),
            "updated": "2026-10-06",
        })
    product_keys = {p["key"] for p in products}

    # Nutrition per 100 g (or 100 ml), with the weight of one unit for products
    # counted in eggs, punnets, cloves and so on.
    with open(NUTRITION_CSV, newline="", encoding="utf-8") as f:
        nutrition = {row["key"]: row for row in csv.DictReader(f)}
    missing = product_keys - nutrition.keys()
    if missing:
        raise SystemExit(f"No nutrition row for: {', '.join(sorted(missing))}")
    for p in products:
        row = nutrition[p["key"]]
        p["nutrition"] = {k: float(row[k]) for k in ("kcal", "protein", "carbs", "fat", "fibre")}
        if p["unit"] not in ("g", "ml"):
            if not row["grams_per_unit"]:
                raise SystemExit(f"{p['key']} is counted in {p['unit']} and needs grams_per_unit")
            p["gramsPerUnit"] = float(row["grams_per_unit"])

    # Recipes (headline info)
    recipes = {}
    order = []
    ws = wb["Recipes"]
    for row in ws.iter_rows(min_row=2, values_only=True):
        number, name, category, _ingredients, _count, notes, link = row[:7]
        if not name:
            continue
        notes = notes or ""
        match = re.search(r"\b((?:Serves|Makes|Fills) [^.]+)\.", notes)
        rid = slug(name)
        recipes[name] = {
            "id": rid,
            "name": name,
            "category": category,
            "notes": notes,
            "yield": match.group(1) if match else "",
            "link": link or "",
            "lines": [],
        }
        order.append(name)

    servings = {}
    if SERVINGS_CSV.exists():
        with open(SERVINGS_CSV, newline="", encoding="utf-8") as f:
            servings = {row["recipe_id"]: float(row["servings"]) for row in csv.DictReader(f)}
    for r in recipes.values():
        r["servings"] = servings.get(r["id"])

    # Ingredient lines
    ws = wb["Ingredient Costing"]
    for row in ws.iter_rows(min_row=2, values_only=True):
        recipe_name, _cat, text, status, key, _prod, qty = row[:7]
        note = row[13] or ""
        if recipe_name not in recipes or not status:
            continue
        recipe = recipes[recipe_name]
        line = {"id": f"{recipe['id']}-{len(recipe['lines']) + 1}", "text": text}
        if status in ("Costed", "Substitute"):
            if key not in product_keys:
                raise SystemExit(f"Unknown product key {key!r} in {recipe_name}")
            line.update(kind="product", productKey=key, qty=num(qty))
            if status == "Substitute":
                line["substitute"] = True
        elif status == "Uses 5 Veg Mince":
            line.update(kind="recipe", recipeId=slug("5 Veg Mince (freezer base)"), qty=num(qty) or 1)
        elif status == "Basic - not costed":
            line["kind"] = "basic"
        elif status == "Not at Woolworths":
            line["kind"] = "missing"
        else:
            raise SystemExit(f"Unknown status {status!r} in {recipe_name}")
        if note:
            line["note"] = note
        recipe["lines"].append(line)

    # Notes & assumptions
    ws = wb["Notes & Assumptions"]
    notes = []
    for row in ws.iter_rows(min_row=5, values_only=True):
        label, text = row[1], row[2]
        if label and text and label != "Price basis (choose)":
            notes.append({"label": label, "text": text})

    excluded = []
    ws = wb["Not included"]
    for row in ws.iter_rows(min_row=2, values_only=True):
        if row[0]:
            excluded.append({"name": row[0], "reason": row[1], "link": row[2] or ""})

    seed = {
        "pricesCaptured": "2026-10-06",
        "notes": notes,
        "excluded": excluded,
        "categories": sorted({r["category"] for r in recipes.values()}),
        "products": products,
        "recipes": [recipes[name] for name in order],
    }
    out.parent.mkdir(parents=True, exist_ok=True)
    out.write_text(json.dumps(seed, ensure_ascii=False, indent=1) + "\n", encoding="utf-8")
    lines = sum(len(r["lines"]) for r in seed["recipes"])
    print(f"Wrote {out}: {len(products)} products, {len(order)} recipes, {lines} ingredient lines")


if __name__ == "__main__":
    main()
